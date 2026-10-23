"""Builders that render structured notes into PDF, DOCX and PPTX bytes.

Used to create the bundled demo document and test fixtures. The optional notice is stored in
the file metadata (subject/comments) rather than the body, so it labels the file without
being analysed as study content.
"""

from __future__ import annotations

import html
import io

Sections = list[tuple[str, list[str]]]


def build_pdf(title: str, sections: Sections, notice: str | None = None) -> bytes:
    import fitz

    parts = [f"<h1>{html.escape(title)}</h1>"]
    for heading, paragraphs in sections:
        parts.append(f"<h2>{html.escape(heading)}</h2>")
        parts.extend(f"<p>{html.escape(p)}</p>" for p in paragraphs)
    css = (
        "body{font-family:sans-serif;font-size:11pt;} h1{font-size:20pt;}"
        " h2{font-size:15pt;margin-top:14pt;} p{margin-bottom:6pt;}"
    )
    story = fitz.Story(html="".join(parts), user_css=css)
    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (56, 56, -56, -56)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    if not notice:
        return buffer.getvalue()
    doc = fitz.open(stream=buffer.getvalue(), filetype="pdf")
    doc.set_metadata({"title": title, "subject": notice, "creator": "StudyForge AI sample generator"})
    return doc.tobytes()


def build_docx(title: str, sections: Sections, notice: str | None = None) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(title, level=0)
    document.core_properties.title = title
    if notice:
        document.core_properties.subject = notice
        document.core_properties.comments = notice
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for p in paragraphs:
            document.add_paragraph(p)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pptx(title: str, sections: Sections, notice: str | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt

    from app.nlp.text import split_sentences

    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    prs.core_properties.title = title
    if notice:
        prs.core_properties.subject = notice
        prs.core_properties.comments = notice
    for heading, paragraphs in sections:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        body = slide.placeholders[1].text_frame
        sentences = [s for p in paragraphs for s in split_sentences(p)]
        body.text = sentences[0]
        for sentence in sentences[1:]:
            para = body.add_paragraph()
            para.text = sentence
        for para in body.paragraphs:
            for run in para.runs:
                run.font.size = Pt(12)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
