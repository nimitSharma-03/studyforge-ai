from __future__ import annotations

import io

import fitz
import pytest

from app.nlp.extract import ExtractionError, detect_file_type, extract
from app.utils.sample_docs import build_docx, build_pdf, build_pptx


def test_pdf_extraction_keeps_headings_and_pages(demo_pdf):
    result = extract("notes.pdf", demo_pdf)
    assert result.file_type == "pdf" and result.unit_label == "page"
    assert result.unit_count >= 2
    headings = [b.text for b in result.blocks if b.kind == "heading"]
    assert "Tokenization" in headings
    assert "Evaluation of NLP Systems" in headings
    assert all(1 <= b.location <= result.unit_count for b in result.blocks)
    # Content near the end of the document must be attributed to a later page than the start.
    last_heading = next(b for b in result.blocks if b.text == "Evaluation of NLP Systems")
    assert last_heading.location > result.blocks[0].location


def test_docx_extraction_uses_heading_styles(demo_docx):
    result = extract("notes.docx", demo_docx)
    assert result.unit_label == "section"
    headings = [b for b in result.blocks if b.kind == "heading"]
    assert headings[0].level == 1  # Title style
    assert any(b.text == "Word Embeddings" for b in headings)
    paragraph = next(b for b in result.blocks if b.text.startswith("Word2Vec is a neural method"))
    heading = next(b for b in headings if b.text == "Word Embeddings")
    assert paragraph.location == heading.location


def test_pptx_extraction_tracks_slide_numbers(demo_pptx):
    result = extract("deck.pptx", demo_pptx)
    assert result.unit_label == "slide"
    assert result.unit_count == 10  # cover + 9 topic slides
    titles = {b.text: b.location for b in result.blocks if b.kind == "heading"}
    assert titles["Tokenization"] == 4
    bullet = next(b for b in result.blocks if b.text.startswith("Tokens are usually words"))
    assert bullet.location == 4


def test_docx_tables_are_extracted():
    import docx

    document = docx.Document()
    document.add_heading("Glossary", level=1)
    document.add_paragraph(
        "The table below lists important terms used throughout this unit of the course notes. "
        "Learn each definition before the mid-semester examination."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Corpus"
    table.cell(0, 1).text = "A large structured collection of texts used to study language"
    table.cell(1, 0).text = "Lemma"
    table.cell(1, 1).text = "The dictionary form of a word used as its canonical representation"
    buffer = io.BytesIO()
    document.save(buffer)
    result = extract("glossary.docx", buffer.getvalue())
    assert any(b.text.startswith("Corpus — A large structured collection") for b in result.blocks)


@pytest.mark.parametrize(
    "name,data,code",
    [
        ("notes.txt", b"plain text file", "unsupported_type"),
        ("notes.exe", b"MZ\x90\x00", "unsupported_type"),
        ("fake.pdf", b"this is not a pdf at all", "corrupted"),
        ("fake.docx", b"%PDF-1.4 not a docx", "corrupted"),
        ("broken.pptx", b"PK\x03\x04garbage-zip", "corrupted"),
    ],
)
def test_invalid_files_are_rejected(name, data, code):
    with pytest.raises(ExtractionError) as info:
        detect_file_type(name, data)
    assert info.value.code == code


def test_docx_with_wrong_internal_structure_is_rejected(demo_pptx):
    # A real PPTX renamed to .docx is a valid zip but not a Word document.
    with pytest.raises(ExtractionError) as info:
        detect_file_type("renamed.docx", demo_pptx)
    assert info.value.code == "corrupted"


def test_empty_pdf_reports_insufficient_text():
    doc = fitz.open()
    doc.new_page()
    with pytest.raises(ExtractionError) as info:
        extract("blank.pdf", doc.tobytes())
    assert info.value.code == "insufficient_text"


def test_scanned_pdf_is_detected():
    doc = fitz.open()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 120), False)
    pix.clear_with(200)
    for _ in range(2):
        page = doc.new_page()
        page.insert_image(fitz.Rect(50, 50, 400, 400), pixmap=pix)
    with pytest.raises(ExtractionError) as info:
        extract("scan.pdf", doc.tobytes())
    assert info.value.code == "scanned_pdf"
    assert "OCR" in info.value.message


def test_short_document_is_rejected():
    data = build_docx("Tiny", [("Heading", ["Too short to study."])])
    with pytest.raises(ExtractionError) as info:
        extract("tiny.docx", data)
    assert info.value.code == "insufficient_text"


def test_repeated_headers_and_page_numbers_are_removed():
    doc = fitz.open()
    body = (
        "Gradient descent is an optimisation algorithm that updates parameters in the direction of the negative "
        "gradient. The learning rate controls the size of each update step during training. "
    )
    for n in range(1, 5):
        page = doc.new_page()
        page.insert_text((72, 40), "CS501 Machine Learning — Lecture Notes", fontsize=9)
        page.insert_textbox(fitz.Rect(72, 80, 520, 700), f"Part {n} of the notes. " + body * 3, fontsize=11)
        page.insert_text((300, 800), str(n), fontsize=9)
    result = extract("lecture.pdf", doc.tobytes())
    texts = [b.text for b in result.blocks]
    assert not any("CS501" in t for t in texts)
    assert not any(t.strip().isdigit() for t in texts)


def test_all_three_formats_produce_same_topics(demo_pdf):
    from app.nlp.pipeline import process
    from app.sample_data.demo_content import NOTICE, SECTIONS, TITLE

    titles = []
    for name, data in (
        ("a.pdf", demo_pdf),
        ("a.docx", build_docx(TITLE, SECTIONS, NOTICE)),
        ("a.pptx", build_pptx(TITLE, SECTIONS, NOTICE)),
    ):
        titles.append([t.title for t in process(extract(name, data)).topics])
    assert titles[0] == titles[1] == titles[2]
    assert titles[0][2] == "Tokenization"


def test_pdf_builder_round_trip():
    data = build_pdf(
        "Doc", [("Section A", ["Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu."] * 4)]
    )
    assert data.startswith(b"%PDF")
